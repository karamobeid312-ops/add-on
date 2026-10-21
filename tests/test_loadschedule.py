# -*- coding: utf-8 -*-
import sample_al_yasat
from sld import loadschedule as ls
from sld.editor import Editor
from sld.model import CircuitInfo, EquipmentInfo, build_schematic


def sample():
    return Editor(build_schematic(*sample_al_yasat.build()))


def sheet(sheets, name):
    return next(s for s in sheets if s.title == name)


def test_one_sheet_per_board_in_tree_order():
    sheets, _, dbs = ls.collect(sample())
    names = [s.title for s in sheets]
    assert names[:2] == ["MDB-1", "SMDB-GF-M1"]
    assert "USMDB-GF-M" in names and "LDB-GF-01" not in names
    assert "LDB-GF-01" in dbs


def test_rows_link_lookup_and_value():
    sheets, by_board, _ = ls.collect(sample())
    rows = dict((r.name, r) for r in sheet(sheets, "SMDB-GF-M1").rows)
    assert rows["SMDB-1ST-01"].source == "link"
    assert rows["LDB-GF-01"].source == "lookup" and rows["LDB-GF-01"].kind == "LTG"
    assert rows["SMDB-1ST-01"].label == "RYB1" and rows["SMDB-1ST-01"].breaker == "250A"
    ups = [r for r in sheet(sheets, "SMDB-GF-M1").rows if r.name == "UPS"]
    assert ups[0].source == "link" and by_board[ups[0].target].title == "USMDB-GF-M"
    assert ups[1].source is None                     # second input way: no load twice
    # a linked board's row carries that board's totals
    child = sheet(sheets, "SMDB-1ST-01").totals()
    assert abs(rows["SMDB-1ST-01"].connected_kw - child[6]) < 1e-6
    assert abs(rows["SMDB-1ST-01"].demand_kw - child[7]) < 1e-6
    ev = dict((r.name, r) for r in sheet(sheets, "MDB-1").rows)["EV-01"]
    assert ev.source == "value" and ev.factor == 1.0


def test_split_phases():
    assert ls.split_phases(9.0) == (3.0, 3.0, 3.0)
    assert ls.split_phases(4.0, "Y") == (0.0, 4.0, 0.0)
    assert ls.split_phases(10.0, typed=(2.0, 2.0, 1.0)) == (4.0, 4.0, 2.0)
    assert ls.split_phases(5.0, phases=1) == (5.0, 0.0, 0.0)
    assert ls.split_phases(None) is None


def test_revit_phase_split_used():
    equipment = [EquipmentInfo("MDB", "MDB-1"),
                 EquipmentInfo("SMDB", "SMDB-1", upstream_protection="250AT/250AF MCCB"),
                 EquipmentInfo("DB", "LDB-1", connected_kw=12.0, demand_kw=9.0,
                               phase_kw=(5.0, 4.0, 3.0))]
    circuits = [CircuitInfo("c1", "MDB", "1", poles="3", start_slot=1, fed_equipment_ids=["SMDB"]),
                CircuitInfo("c2", "SMDB", "1", poles="3", start_slot=1, rating="63 A",
                            fed_equipment_ids=["DB"],
                            cable=u"(4X16)mm² CU/XLPE/PVC +(1X16)mm² CU/PVC(E)")]
    sheets, _, dbs = ls.collect(Editor(build_schematic(equipment, circuits)))
    row = sheet(sheets, "SMDB-1").rows[0]
    assert row.phase_kw == (5.0, 4.0, 3.0) and abs(row.factor - 0.75) < 1e-9
    assert dbs["LDB-1"][0] == 12.0
    assert ls.incoming_text(row.cable) == u"4Cx16mm² CU/XLPE/PVC + 1Cx16mm² ECC"


def test_load_types():
    assert ls.load_type("LDB-GF-01") == "LTG"
    assert ls.load_type("ADB-GF-02") == "AC"
    assert ls.load_type("FCU-01") == "AC"
    assert ls.load_type("PDB-GF-01") == "PWR"
    assert ls.load_type("LIFT-01") == "PWR"
    assert ls.load_type("WATER HEATER") == "WH"
    assert ls.load_type("PUMP-01") == "MO"


def test_sheet_names_unique_and_valid():
    used = set()
    assert ls.sheet_name("A/B", used) == "A-B"
    assert ls.sheet_name("a/b", used) == "a-b (2)"
    assert len(ls.sheet_name("X" * 40, used)) == 31


def test_workbook_formulas(tmp_path):
    wb, sheets = ls.build(sample(), {"project": "KEZAD", "location": "ABU DHABI"})
    path = str(tmp_path / "schedules.xlsx")
    wb.save(path)
    import openpyxl
    book = openpyxl.load_workbook(path)
    assert book.sheetnames[0] == "MDB-1" and book.sheetnames[-1] == ls.SCHEDULE_SHEET
    ws = book["SMDB-GF-M1"]
    assert ws["B3"].value == "SMDB-GF-M1" and ws["B4"].value == "MDB-1"
    assert ws["K8"].value == "='SMDB-1ST-01'!K18"           # board row: its sheet's totals
    assert ws["N10"].value.startswith("=VLOOKUP(B10,")      # DB row: Revit panel loads
    t = sheet(sheets, "SMDB-GF-M1").total_row
    assert ws["U%d" % (t + 1)].value == "=SUM(N8:N%d)" % (t - 1)
    assert "Project/Building Name : KEZAD" in ws["A1"].value
